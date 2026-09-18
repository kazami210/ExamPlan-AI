import io
from docx import Document
from backend.services.parsers import parse_docx, parse_txt

# Test DOCX
doc = Document()
doc.add_heading("ĐỀ CƯƠNG KINH TẾ VI MÔ", 0)
doc.add_paragraph("Chương 1: Tổng quan về kinh tế học vi mô")
doc.add_paragraph("Chương 2: Cung, cầu và giá cả thị trường")
doc_stream = io.BytesIO()
doc.save(doc_stream)
doc_bytes = doc_stream.getvalue()

docx_text = parse_docx(doc_bytes)
assert "KINH TẾ VI MÔ" in docx_text
assert "Chương 1" in docx_text
print("✓ DOCX parser test passed!")

# Test TXT
txt_bytes = "Chương 1: Khái niệm cơ bản\nChương 2: Thuật toán".encode("utf-8")
txt_text = parse_txt(txt_bytes)
assert "Thuật toán" in txt_text
print("✓ TXT parser test passed!")
